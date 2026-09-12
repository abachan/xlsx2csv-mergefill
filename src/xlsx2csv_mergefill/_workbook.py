"""Excel ワークブックをセルデータ読取用に安全に読み込む内部処理。"""

from __future__ import annotations

import io
import re
import xml.etree.ElementTree as ET
import zipfile
from contextlib import contextmanager
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from ._filename_formula import _refresh_filename_formula_caches


_DRAWING_RELATIONSHIP_SUFFIX = "/drawing"
_SPREADSHEETML_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_PHONETIC_ATTRIBUTE_PATTERN = re.compile(r'\s+phonetic="[^"]*"')
_PHONETIC_PROPERTY_PATTERN = re.compile(
    r"<phoneticPr\b[^>]*(?:/>|>.*?</phoneticPr>)",
    flags=re.DOTALL,
)
_PHONETIC_RUN_PATTERN = re.compile(r"<rPh\b[^>]*>.*?</rPh>", flags=re.DOTALL)


def _strip_drawing_relationships(content: bytes) -> bytes:
    """ワークシート等から図・画像・グラフへの関連だけを除去する。"""
    if b"/drawing" not in content:
        return content

    root = ET.fromstring(content)
    removed = False
    for relationship in list(root):
        relationship_type = relationship.get("Type", "")
        if relationship_type.endswith(_DRAWING_RELATIONSHIP_SUFFIX):
            root.remove(relationship)
            removed = True

    if not removed:
        return content
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _repair_chartsheet_custom_views(content: bytes) -> bytes:
    """openpyxl が読み込めない、scale 未指定のグラフシート表示設定を補正する。"""
    if b"customSheetView" not in content:
        return content

    root = ET.fromstring(content)
    changed = False
    custom_view_tag = f"{{{_SPREADSHEETML_NAMESPACE}}}customSheetView"
    for custom_view in root.iter(custom_view_tag):
        if custom_view.get("scale") is None:
            # OOXML 上の既定表示倍率。CSV化するセル値には影響しない。
            custom_view.set("scale", "100")
            changed = True

    if not changed:
        return content
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _prepare_xlsx_for_cell_data(
    data: io.BytesIO,
    workbook_filename: str | None = None,
) -> io.BytesIO:
    """セル値の取得を妨げる不要な OOXML 要素・属性を除去する。

    除去対象:
    - <col> 要素の phonetic 属性（ColumnDimension TypeError の原因）
    - <phoneticPr> 要素（列・シートのふりがな表示設定）
    - <rPh> 要素（セル内リッチテキスト中のルビテキスト）
    - ワークシート・グラフシートの drawing 関連（画像・グラフ・図形）
    - 安全に評価できるブック名依存式の古い保存済み値

    グラフシートの customSheetView に scale がない場合は、OOXML の既定値で
    補正する。入力ファイル自体は変更せず、openpyxl に渡すメモリ上のコピー
    だけを処理する。
    """
    with zipfile.ZipFile(data, "r") as source:
        first_change = _find_first_changed_part(source, workbook_filename)
        if first_change is None:
            data.seek(0)
            return data

        output = io.BytesIO()
        try:
            with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
                first_item, first_content = first_change
                for item in source.infolist():
                    if item is first_item:
                        content = first_content
                    else:
                        content = source.read(item)
                        content = _transform_xlsx_part(
                            item.filename,
                            content,
                            workbook_filename,
                        )
                    target.writestr(item, content)
        except Exception:
            output.close()
            raise

    output.seek(0)
    return output


def _find_first_changed_part(
    source: zipfile.ZipFile,
    workbook_filename: str | None,
) -> tuple[zipfile.ZipInfo, bytes] | None:
    """書き換えが必要な最初のOOXML部品と変換済み内容を返す。"""
    for item in source.infolist():
        if not _is_transform_candidate(item.filename):
            continue

        content = source.read(item)
        transformed = _transform_xlsx_part(
            item.filename,
            content,
            workbook_filename,
        )
        if transformed != content:
            return item, transformed
    return None


def _is_transform_candidate(filename: str) -> bool:
    """セル値読取のために確認が必要なOOXML部品かを返す。"""
    if filename.startswith("xl/worksheets/") and filename.endswith(".xml"):
        return True
    if (
        filename.startswith("xl/worksheets/_rels/")
        or filename.startswith("xl/chartsheets/_rels/")
    ) and filename.endswith(".rels"):
        return True
    return filename.startswith("xl/chartsheets/") and filename.endswith(".xml")


def _transform_xlsx_part(
    filename: str,
    content: bytes,
    workbook_filename: str | None,
) -> bytes:
    """OOXML部品へ必要な変換だけを逐次適用する。"""
    if filename.startswith("xl/worksheets/") and filename.endswith(".xml"):
        content = _strip_phonetic_xml(content)
        if workbook_filename is not None:
            content = _refresh_filename_formula_caches(content, workbook_filename)
        return content
    if (
        filename.startswith("xl/worksheets/_rels/")
        or filename.startswith("xl/chartsheets/_rels/")
    ) and filename.endswith(".rels"):
        return _strip_drawing_relationships(content)
    if filename.startswith("xl/chartsheets/") and filename.endswith(".xml"):
        return _repair_chartsheet_custom_views(content)
    return content


def _strip_phonetic_from_xlsx(data: io.BytesIO) -> io.BytesIO:
    """互換性のために残す、旧名称の前処理関数。"""
    return _prepare_xlsx_for_cell_data(data)


def _strip_phonetic_xml(content: bytes) -> bytes:
    """ワークシート XML からルビ関連の要素・属性を除去する。"""
    # UTF-16 BOM を考慮してデコード（OOXML は UTF-8 か UTF-16 のみ合法）
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        xml_text = content.decode("utf-16")
    else:
        if not any(
            marker in content for marker in (b"phonetic=", b"<phoneticPr", b"<rPh")
        ):
            return content
        xml_text = content.decode("utf-8")

    stripped_xml = _PHONETIC_ATTRIBUTE_PATTERN.sub("", xml_text)
    stripped_xml = _PHONETIC_PROPERTY_PATTERN.sub("", stripped_xml)
    stripped_xml = _PHONETIC_RUN_PATTERN.sub("", stripped_xml)
    if stripped_xml == xml_text:
        return content
    return stripped_xml.encode("utf-8")


def _load_workbook(input_xlsx: Path | str) -> Workbook:
    """Excelワークブックを読み込む共通関数。"""
    input_path = Path(input_xlsx)
    if not input_path.exists():
        raise FileNotFoundError(f"入力ファイルが見つかりません: {input_path}")

    # BytesIO 経由にすることで解析中に例外が発生しても、Windows 上で
    # 入力ファイルのロックが残らないようにする。
    with input_path.open("rb") as input_file:
        source_data = io.BytesIO(input_file.read())

    prepared_data: io.BytesIO | None = None
    try:
        prepared_data = _prepare_xlsx_for_cell_data(source_data, input_path.name)
        return load_workbook(
            filename=prepared_data,
            data_only=True,
            read_only=False,
        )
    finally:
        if prepared_data is not None and prepared_data is not source_data:
            prepared_data.close()
        source_data.close()


@contextmanager
def _open_workbook(
    input_xlsx: Path | str,
) -> Iterator[Workbook]:
    """ワークブックを開き、利用後に確実に閉じる。"""
    workbook = _load_workbook(input_xlsx)
    try:
        yield workbook
    finally:
        workbook.close()
