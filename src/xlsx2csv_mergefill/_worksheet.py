"""ワークシートから表示対象のセル値を抽出する内部処理。"""

from __future__ import annotations

from typing import Dict, Iterable, List, Optional, Tuple

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.worksheet.worksheet import Worksheet


CellCoord = Tuple[int, int]

# UTF-8換算で101バイト以上の値は結合セル範囲へ補填しない
DEFAULT_MERGE_FILL_MAX_BYTES = 100


def _get_display_value(cell: Cell | MergedCell | None) -> Optional[object]:
    """Excel 上で実際に見える値だけを返す。"""
    if cell is None or isinstance(cell, MergedCell):
        return None

    value = cell.value
    if value is None:
        return None

    hyperlink = getattr(cell, "hyperlink", None)
    if hyperlink is not None:
        location = getattr(hyperlink, "location", None)
        target = getattr(hyperlink, "target", None)
        # hyperlink 要素のみで定義された内部リンクは、openpyxl が参照先を
        # cell.value に設定することがある。Excel 上ではセル値として表示されないため除外する。
        if target is None and location and value == location:
            return None

    return value


def _should_merge_fill(value: Optional[object], merge_fill_max_bytes: int) -> bool:
    """マージセル範囲へ値を補填してよいかを返す。"""
    return value is None or len(str(value).encode("utf-8")) <= merge_fill_max_bytes


def _build_merged_value_map(
    ws: Worksheet,
    merge_fill_max_bytes: int = DEFAULT_MERGE_FILL_MAX_BYTES,
) -> Dict[CellCoord, Optional[object]]:
    """マージセルの値マップを構築する。"""
    merged_map: Dict[CellCoord, Optional[object]] = {}
    for merged_range in ws.merged_cells.ranges:
        # ws.cell() はセルを生成・キャッシュするため、_cells.get() で既存セルのみ参照する。
        top_left = ws._cells.get((merged_range.min_row, merged_range.min_col))
        top_left_value = _get_display_value(top_left)
        if not _should_merge_fill(top_left_value, merge_fill_max_bytes):
            continue

        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for column in range(merged_range.min_col, merged_range.max_col + 1):
                merged_map[(row, column)] = top_left_value
    return merged_map


def _get_cell_value(
    ws: Worksheet,
    merged_map: Dict[CellCoord, Optional[object]],
    row: int,
    column: int,
) -> Optional[object]:
    """マージセルを考慮してセル値を返す。"""
    key = (row, column)
    if key in merged_map:
        return merged_map[key]
    cell = ws._cells.get(key)
    return _get_display_value(cell)


def _effective_bounds(
    ws: Worksheet,
    merged_map: Dict[CellCoord, Optional[object]],
) -> Tuple[int, int]:
    """値（None以外）が存在する最終行・最終列を返す。"""
    last_row = last_column = 0
    for (row, column), cell in ws._cells.items():
        value = merged_map.get((row, column), _get_display_value(cell))
        if value is not None:
            last_row = max(last_row, row)
            last_column = max(last_column, column)

    # マージ範囲内の非左上隅セルも考慮する。
    for (row, column), value in merged_map.items():
        if value is not None:
            last_row = max(last_row, row)
            last_column = max(last_column, column)
    return last_row, last_column


def _iter_rows_values(
    ws: Worksheet,
    merge_fill_max_bytes: int = DEFAULT_MERGE_FILL_MAX_BYTES,
) -> Iterable[List[Optional[object]]]:
    """ワークシートからマージセル展開済みの行データを取得する。"""
    merged_map = _build_merged_value_map(ws, merge_fill_max_bytes)
    last_row, last_column = _effective_bounds(ws, merged_map)
    for row in range(1, last_row + 1):
        yield [
            _get_cell_value(ws, merged_map, row, column)
            for column in range(1, last_column + 1)
        ]
