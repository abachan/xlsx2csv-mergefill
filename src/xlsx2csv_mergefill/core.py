"""xlsx2csv_mergefill の公開APIと変換処理のオーケストレーション。"""

from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Optional

from openpyxl.worksheet.worksheet import Worksheet

# 従来 core.py に存在した内部名も、互換性のためこのモジュールから参照可能にする。
from ._csv import (  # noqa: F401 -- 旧内部名を互換性のため再公開する
    _sanitize_filename,
    _serialize_row,
    _write_csv,
    to_csv_string,
)
from ._workbook import (  # noqa: F401 -- 旧内部名を互換性のため再公開する
    _load_workbook,
    _open_workbook,
    _prepare_xlsx_for_cell_data,
    _repair_chartsheet_custom_views,
    _strip_drawing_relationships,
    _strip_phonetic_from_xlsx,
)
from ._worksheet import (  # noqa: F401 -- 旧内部名を互換性のため再公開する
    DEFAULT_MERGE_FILL_MAX_BYTES,
    CellCoord,
    _build_merged_value_map,
    _effective_bounds,
    _get_cell_value,
    _get_display_value,
    _iter_rows_values,
    _should_merge_fill,
)


def convert_file(
    input_xlsx: Path | str,
    output_csv: Path | str,
    use_numeric_sheet_names: bool = False,
    include_hidden_sheets: bool = False,
    merge_fill_max_bytes: int = DEFAULT_MERGE_FILL_MAX_BYTES,
) -> int:
    """Excelファイルをシート単位のCSVファイルへ変換する。

    Args:
        input_xlsx: 入力Excelファイルのパス
        output_csv: 出力CSVファイルのパス
        use_numeric_sheet_names: Trueの場合、シート名を数値でカウントアップ（0, 1, 2...）
        include_hidden_sheets: Trueの場合、非表示シートも出力する。デフォルトはFalse
        merge_fill_max_bytes: マージセル補填する値の最大バイト数（UTF-8換算）。デフォルトは100
    """
    with _open_workbook(input_xlsx) as workbook:
        # グラフシートを含め、Excel 上のシート順をそのまま使用する。
        # グラフシート自体にはセルがないため、空の CSV として出力する。
        sheets = [workbook[sheet_name] for sheet_name in workbook.sheetnames]

        if not include_hidden_sheets:
            sheets = [sheet for sheet in sheets if sheet.sheet_state == "visible"]

        if not sheets:
            raise ValueError("ワークシートが見つかりません")

        output_path = Path(output_csv)
        base_output = output_path.with_suffix("")

        for index, sheet in enumerate(sheets):
            sheet_identifier = (
                str(index)
                if use_numeric_sheet_names
                else _sanitize_filename(sheet.title)
            )
            target_path = (
                base_output.parent / f"{base_output.name}_{sheet_identifier}.csv"
            )
            rows = (
                _iter_rows_values(sheet, merge_fill_max_bytes)
                if isinstance(sheet, Worksheet)
                else []
            )
            _write_csv(rows, target_path)

        return 0


# Backward-compatible alias (deprecated)
def excel_to_csv(
    input_xlsx: Path | str,
    output_csv: Path | str,
    use_numeric_sheet_names: bool = False,
    include_hidden_sheets: bool = False,
    merge_fill_max_bytes: int = DEFAULT_MERGE_FILL_MAX_BYTES,
) -> int:
    """Deprecated alias for convert_file. Will be removed in a future release."""
    return convert_file(
        input_xlsx,
        output_csv,
        use_numeric_sheet_names,
        include_hidden_sheets,
        merge_fill_max_bytes,
    )


def read_sheet(
    input_xlsx: Path | str,
    merge_fill_max_bytes: int = DEFAULT_MERGE_FILL_MAX_BYTES,
) -> List[List[Optional[object]]]:
    """Excelファイルのアクティブなワークシートをリスト形式で返す。"""
    with _open_workbook(input_xlsx) as workbook:
        worksheet = workbook.active
        if not isinstance(worksheet, Worksheet):
            if not workbook.worksheets:
                raise ValueError("ワークシートが見つかりません")
            worksheet = workbook.worksheets[0]
        return list(_iter_rows_values(worksheet, merge_fill_max_bytes))


# Backward-compatible alias (deprecated)
def load_excel_data(
    input_xlsx: Path | str,
    merge_fill_max_bytes: int = DEFAULT_MERGE_FILL_MAX_BYTES,
) -> List[List[Optional[object]]]:
    """Deprecated alias for read_sheet. Will be removed in a future release."""
    return read_sheet(input_xlsx, merge_fill_max_bytes)


def read_workbook(
    input_xlsx: Path | str,
    merge_fill_max_bytes: int = DEFAULT_MERGE_FILL_MAX_BYTES,
) -> Dict[str, List[List[Optional[object]]]]:
    """Excelファイルの全シートを辞書形式で返す。"""
    with _open_workbook(input_xlsx) as workbook:
        return {
            sheet.title: (
                list(_iter_rows_values(sheet, merge_fill_max_bytes))
                if isinstance(sheet, Worksheet)
                else []
            )
            for sheet in (workbook[sheet_name] for sheet_name in workbook.sheetnames)
        }


# Backward-compatible alias (deprecated)
def load_all_sheets_data(
    input_xlsx: Path | str,
    merge_fill_max_bytes: int = DEFAULT_MERGE_FILL_MAX_BYTES,
) -> Dict[str, List[List[Optional[object]]]]:
    """Deprecated alias for read_workbook. Will be removed in a future release."""
    return read_workbook(input_xlsx, merge_fill_max_bytes)


def list_sheets(input_xlsx: Path | str) -> List[str]:
    """Excelファイルのシート名一覧を取得する。"""
    with _open_workbook(input_xlsx) as workbook:
        return workbook.sheetnames


# Backward-compatible alias (deprecated)
def get_sheet_names(input_xlsx: Path | str) -> List[str]:
    """Deprecated alias for list_sheets. Will be removed in a future release."""
    return list_sheets(input_xlsx)


# Backward-compatible alias (deprecated)
def data_to_csv_string(data: List[List[Optional[object]]]) -> str:
    """Deprecated alias for to_csv_string. Will be removed in a future release."""
    return to_csv_string(data)
